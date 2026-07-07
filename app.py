import io
import json
import os
import sqlite3
import time
from datetime import datetime
from uuid import uuid4

import cv2
import numpy as np
from flask import (
    Flask,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename
from ultralytics import YOLO

# ---------- Конфигурация ----------
BASE_DIR = os.path.abspath(os.path.dirname(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, "static", "uploads")
RESULT_FOLDER = os.path.join(BASE_DIR, "static", "results")
DB_PATH = os.path.join(BASE_DIR, "segmentation_results.db")
MODEL_PATH = os.path.join(BASE_DIR, "yolov8n-seg.pt")

ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "bmp", "tiff", "tif"}
MAX_FILE_SIZE = 16 * 1024 * 1024

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["RESULT_FOLDER"] = RESULT_FOLDER
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

if not os.path.isfile(MODEL_PATH):
    raise FileNotFoundError(
        f"Файл модели не найден: {MODEL_PATH}. "
        "Поместите yolov8n-seg.pt рядом с app.py."
    )

# Модель загружается один раз при старте приложения.
model = YOLO(MODEL_PATH)


# ---------- Работа с БД ----------
def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_database():
    with get_db_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS segmentation_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                image_name TEXT NOT NULL,
                image_path TEXT NOT NULL,
                processed_at TEXT NOT NULL,
                num_objects INTEGER,
                total_time REAL,
                boxes_json TEXT,
                confidences_json TEXT,
                masks_json TEXT,
                model_used TEXT
            )
            """
        )
        conn.commit()


def save_result_to_db(
    image_name,
    image_path,
    num_objects,
    total_time,
    boxes,
    confidences,
    contours_list,
    model_name="yolov8n-seg.pt",
):
    boxes_json = json.dumps(
        boxes.tolist() if boxes is not None else None,
        ensure_ascii=False,
    )
    confidences_json = json.dumps(
        confidences.tolist() if confidences is not None else None,
        ensure_ascii=False,
    )
    contours_json = json.dumps(
        contours_list if contours_list else None,
        ensure_ascii=False,
    )

    with get_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO segmentation_results
            (image_name, image_path, processed_at, num_objects, total_time,
             boxes_json, confidences_json, masks_json, model_used)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                image_name,
                image_path,
                datetime.now().isoformat(timespec="seconds"),
                num_objects,
                total_time,
                boxes_json,
                confidences_json,
                contours_json,
                model_name,
            ),
        )
        conn.commit()


def get_all_results():
    with get_db_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM segmentation_results ORDER BY id DESC"
        ).fetchall()
    return rows


def rows_to_dicts(rows):
    data = []
    for row in rows:
        data.append(
            {
                "id": row["id"],
                "image_name": row["image_name"],
                "image_path": row["image_path"],
                "processed_at": row["processed_at"],
                "num_objects": row["num_objects"],
                "total_time": row["total_time"],
                "boxes": json.loads(row["boxes_json"]) if row["boxes_json"] else None,
                "confidences": (
                    json.loads(row["confidences_json"])
                    if row["confidences_json"]
                    else None
                ),
                "contours": json.loads(row["masks_json"]) if row["masks_json"] else None,
                "model": row["model_used"],
            }
        )
    return data


# ---------- Работа с изображениями ----------
def allowed_file(filename):
    return (
        bool(filename)
        and "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def extract_contours_from_masks(masks):
    """Преобразует маски модели в списки точек внешних контуров."""
    if masks is None:
        return []

    if hasattr(masks, "data"):
        masks_np = masks.data.cpu().numpy()
    else:
        masks_np = np.asarray(masks)

    contours_list = []
    for mask in masks_np:
        binary_mask = (mask > 0.5).astype(np.uint8) * 255
        contours, _ = cv2.findContours(
            binary_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        if contours:
            contour = max(contours, key=cv2.contourArea)
            points = contour.squeeze().tolist()
            if isinstance(points, list) and points:
                contours_list.append(points)

    return contours_list


def get_class_name(class_id):
    """Возвращает реальное имя класса из словаря классов загруженной модели."""
    names = model.names
    if isinstance(names, dict):
        return names.get(int(class_id), str(class_id))
    return names[int(class_id)] if int(class_id) < len(names) else str(class_id)


def draw_contours_on_image(
    image, masks, boxes=None, class_ids=None, confidences=None
):
    result_img = image.copy()
    if masks is None:
        return result_img

    if hasattr(masks, "data"):
        masks_np = masks.data.cpu().numpy()
    else:
        masks_np = np.asarray(masks)

    for i, mask in enumerate(masks_np):
        binary_mask = (mask > 0.5).astype(np.uint8) * 255
        contours, _ = cv2.findContours(
            binary_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        cv2.drawContours(result_img, contours, -1, (0, 255, 0), 2)

        if boxes is not None and i < len(boxes):
            x1, y1, x2, y2 = map(int, boxes[i])
            cv2.rectangle(result_img, (x1, y1), (x2, y2), (0, 255, 0), 1)

            if class_ids is not None and confidences is not None and i < len(confidences):
                class_name = get_class_name(class_ids[i])
                label = f"{class_name} {confidences[i]:.2f}"
                cv2.putText(
                    result_img,
                    label,
                    (x1, max(15, y1 - 5)),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.5,
                    (0, 255, 0),
                    1,
                )

    return result_img


def process_single_image(image_path, output_path):
    """Обрабатывает одно изображение и возвращает метаданные результата."""
    image = cv2.imread(image_path)
    if image is None:
        return None

    image_rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)

    start_time = time.perf_counter()
    results = model.predict(
        source=image_rgb,
        imgsz=640,
        conf=0.25,
        verbose=False,
    )
    inference_time = time.perf_counter() - start_time

    boxes = None
    confidences = None
    class_ids = None
    contours_list = []
    masks = None

    if results:
        result = results[0]
        masks = result.masks
        if result.boxes is not None:
            boxes = result.boxes.xyxy.cpu().numpy()
            confidences = result.boxes.conf.cpu().numpy()
            class_ids = result.boxes.cls.cpu().numpy().astype(int)
        contours_list = extract_contours_from_masks(masks)

    num_objects = len(boxes) if boxes is not None else 0
    result_image = draw_contours_on_image(
        image,
        masks,
        boxes=boxes,
        class_ids=class_ids,
        confidences=confidences,
    )

    write_ok = cv2.imwrite(output_path, result_image)
    if not write_ok:
        raise OSError(f"Не удалось сохранить результат: {output_path}")

    detected_classes = []
    if class_ids is not None:
        detected_classes = [get_class_name(class_id) for class_id in class_ids]

    return {
        "num_objects": num_objects,
        "total_time": inference_time,
        "boxes": boxes,
        "confidences": confidences,
        "contours_list": contours_list,
        "classes": detected_classes,
    }


# Инициализируем БД независимо от способа запуска приложения
init_database()


# ---------- Экспорт отчёта ----------
def build_excel_report():
    rows = get_all_results()
    results = rows_to_dicts(rows)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    summary.append(["Отчёт о результатах сегментации МРТ-изображений"])
    summary["A1"].font = Font(size=14, bold=True)
    summary.append([])
    summary.append(["Количество обработанных изображений", len(results)])

    total_objects = sum((r["num_objects"] or 0) for r in results)
    summary.append(["Общее количество обнаруженных объектов", total_objects])

    times = [r["total_time"] for r in results if r["total_time"] is not None]
    average_time = sum(times) / len(times) if times else 0
    summary.append(["Среднее время обработки, с", round(average_time, 3)])
    summary.append(["Модель", results[0]["model"] if results else os.path.basename(MODEL_PATH)])
    summary.column_dimensions["A"].width = 45
    summary.column_dimensions["B"].width = 30

    sheet = workbook.create_sheet("Результаты")
    headers = [
        "ID",
        "Имя изображения",
        "Дата и время",
        "Количество объектов",
        "Время обработки, с",
        "Модель",
        "Путь к исходному файлу",
        "Ограничивающие рамки",
        "Уверенности",
        "Контуры сегментации",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for result in results:
        sheet.append(
            [
                result["id"],
                result["image_name"],
                result["processed_at"],
                result["num_objects"],
                round(result["total_time"], 3) if result["total_time"] is not None else None,
                result["model"],
                result["image_path"],
                json.dumps(result["boxes"], ensure_ascii=False),
                json.dumps(result["confidences"], ensure_ascii=False),
                json.dumps(result["contours"], ensure_ascii=False),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 28, 22, 20, 20, 22, 45, 50, 40, 60]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


# ---------- Маршруты Flask ----------
@app.route("/")
def index():
    results = get_all_results()
    return render_template("index.html", results=results)


@app.route("/upload", methods=["POST"])
def upload_files():
    """Загружает один или несколько файлов, обрабатывает их и сохраняет результаты."""
    files = request.files.getlist("files")
    if not files or all(not file.filename for file in files):
        return redirect(url_for("index"))

    for file in files:
        if not file or not allowed_file(file.filename):
            continue

        original_filename = secure_filename(file.filename)
        unique_prefix = uuid4().hex[:8]
        filename = f"{unique_prefix}_{original_filename}"
        input_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)

        file.save(input_path)

        name, ext = os.path.splitext(filename)
        output_filename = f"{name}_result{ext}"
        output_path = os.path.join(app.config["RESULT_FOLDER"], output_filename)

        try:
            data = process_single_image(input_path, output_path)
            if data is None:
                continue

            save_result_to_db(
                image_name=filename,
                image_path=input_path,
                num_objects=data["num_objects"],
                total_time=data["total_time"],
                boxes=data["boxes"],
                confidences=data["confidences"],
                contours_list=data["contours_list"],
            )
        except Exception:
            # Не оставляем в системе битый результат при ошибке обработки.
            if os.path.exists(input_path):
                os.remove(input_path)
            if os.path.exists(output_path):
                os.remove(output_path)
            raise

    return redirect(url_for("index"))


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


@app.route("/results/<path:filename>")
def result_file(filename):
    return send_from_directory(app.config["RESULT_FOLDER"], filename)


@app.route("/api/results")
def api_results():
    """JSON-эндпоинт для получения истории обработки из БД."""
    return jsonify(rows_to_dicts(get_all_results()))


@app.route("/export/xlsx")
def export_xlsx():
    """Формирует и скачивает отдельный Excel-отчёт по истории обработки."""
    report = build_excel_report()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"segmentation_report_{timestamp}.xlsx"
    return send_file(
        report,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.errorhandler(413)
def request_entity_too_large(_error):
    return (
        jsonify({"error": "Размер загружаемого файла превышает допустимый предел 16 МБ."}),
        413,
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
